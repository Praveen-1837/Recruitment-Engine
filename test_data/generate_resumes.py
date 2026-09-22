import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

os.makedirs("test_data/resumes", exist_ok=True)

# 1. Generate alex_senior_python.pdf
def create_alex_pdf(filename: str):
    c = canvas.Canvas(filename, pagesize=letter)
    width, height = letter
    y = height - 50

    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, y, "Alex Chen - Senior Python Engineer")
    y -= 25

    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Email: alex.chen@example.com | Phone: +1-555-0199 | Location: San Francisco, CA")
    y -= 30

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "PROFESSIONAL SUMMARY")
    y -= 15
    c.setFont("Helvetica", 10)
    summary = (
        "Seasoned Senior Python Engineer with 7 years of hands-on experience designing and building high-scale "
        "distributed backend systems and microservices. Expert in Python, FastAPI, PostgreSQL, and AWS cloud infrastructure."
    )
    c.drawString(50, y, summary[:95])
    y -= 15
    c.drawString(50, y, summary[95:])
    y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "TECHNICAL SKILLS")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "• Programming Languages: Python, SQL, Go, Bash")
    y -= 15
    c.drawString(50, y, "• Frameworks & Libraries: FastAPI, SQLAlchemy, Pydantic, Celery, Redis")
    y -= 15
    c.drawString(50, y, "• Databases: PostgreSQL, MongoDB, Redis")
    y -= 15
    c.drawString(50, y, "• Cloud & DevOps: AWS (EC2, S3, RDS, Lambda), Docker, Kubernetes, CI/CD, Linux, Git")
    y -= 15
    c.drawString(50, y, "• Soft Skills: Engineering Leadership, Mentoring, Agile/Scrum, System Design")
    y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "WORK EXPERIENCE")
    y -= 15
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Staff Backend Engineer | Nexus Cloud Inc. (2021 - Present | 3 years)")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "- Architected FastAPI microservices serving 40k req/sec with sub-50ms latency.")
    y -= 15
    c.drawString(50, y, "- Managed PostgreSQL clusters with database partitioning and query optimization.")
    y -= 15
    c.drawString(50, y, "- Led team of 6 engineers, spearheading containerization via Docker and Kubernetes on AWS.")
    y -= 20

    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Senior Software Engineer | FinTech Systems (2019 - 2021 | 2 years)")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "- Developed REST API endpoints using Python, FastAPI, and SQLAlchemy.")
    y -= 15
    c.drawString(50, y, "- Built automated CI/CD deployment pipelines using GitHub Actions and AWS.")
    y -= 20

    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Software Engineer | DataFlow Labs (2017 - 2019 | 2 years)")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "- Built data ingest pipelines in Python and PostgreSQL.")
    y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "EDUCATION & CERTIFICATIONS")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "• B.S. in Computer Science - University of California, Berkeley")
    y -= 15
    c.drawString(50, y, "• AWS Certified Solutions Architect - Associate")
    y -= 15
    c.drawString(50, y, "• CKA (Certified Kubernetes Administrator)")

    c.save()

# 2. Generate sarah_mid_python.docx
def create_sarah_docx(filename: str):
    doc = Document()
    doc.add_heading("Sarah Jenkins", level=0)
    doc.add_paragraph("Email: sarah.jenkins@example.com | Phone: +1-555-0144 | Location: Austin, TX")
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Mid-level Python Developer with 4 years of experience building reliable web applications and REST APIs. "
        "Proficient in Python, Django, PostgreSQL, and Git with strong collaborative teamwork and agile practices."
    )

    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Languages: Python, JavaScript, SQL")
    doc.add_paragraph("Frameworks: Django, Django REST Framework, Flask")
    doc.add_paragraph("Databases: PostgreSQL, MySQL")
    doc.add_paragraph("Tools: Git, Linux, Docker, Postman, JIRA")
    doc.add_paragraph("Soft Skills: Teamwork, Problem Solving, Communication, Agile/Scrum")

    doc.add_heading("Professional Experience", level=1)
    doc.add_paragraph("Software Engineer | Horizon Digital (2022 - Present | 2 years)")
    doc.add_paragraph("- Built scalable REST APIs using Django and PostgreSQL.")
    doc.add_paragraph("- Integrated third-party payment gateways and webhook listeners.")
    doc.add_paragraph("- Collaborated in cross-functional agile teams.")

    doc.add_paragraph("Associate Developer | WebSphere Solutions (2020 - 2022 | 2 years)")
    doc.add_paragraph("- Maintained Django web portals and optimized SQL database queries.")
    doc.add_paragraph("- Created automated unit tests achieving 85% code coverage.")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Software Engineering - University of Texas at Austin")

    doc.save(filename)

# 3. Generate jordan_junior_python.txt
def create_jordan_txt(filename: str):
    content = """Jordan Miller - Junior Software Developer
Email: jordan.miller@example.com | Phone: +1-555-0188 | Location: Chicago, IL

SUMMARY:
Recent Computer Science graduate with 1 year of professional internship and junior developer experience.
Enthusiastic about backend web development, clean code, and fast learning.

CORE SKILLS:
- Technical Skills: Python, Flask, SQLite, HTML5, CSS3, JavaScript, REST API
- Tools: Git, GitHub, VS Code, Linux
- Soft Skills: Communication, Curiosity, Problem Solving, Teamwork
- Languages: English

WORK EXPERIENCE:
Junior Backend Developer | CloudSpark (2023 - Present | 1 year of experience)
- Developed lightweight web services and internal endpoints using Python and Flask.
- Designed database schemas and queries using SQLite and SQLAlchemy.
- Participated in daily standups, code reviews, and Git feature branching.

Software Development Intern | TechVentures (2022 - 2023 | 6 months)
- Assisted backend engineering team with bug fixes, test scripts, and API documentation.

EDUCATION:
B.S. in Computer Science - University of Illinois (2023)
"""
    with open(filename, "w", encoding="utf-8") as f:
        f.write(content)

# 4. Generate emily_frontend_lead.pdf
def create_emily_pdf(filename: str):
    c = canvas.Canvas(filename, pagesize=letter)
    width, height = letter
    y = height - 50

    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, y, "Emily Zhao - Frontend Engineering Lead")
    y -= 25

    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Email: emily.zhao@example.com | Location: Seattle, WA")
    y -= 30

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "PROFESSIONAL SUMMARY")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Lead Frontend Engineer with 6 years of experience specializing in React, TypeScript, Next.js,")
    y -= 15
    c.drawString(50, y, "and modern UI/UX design systems. Basic Python scripting experience (1 year).")
    y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "SKILLS")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "• Frontend: React, TypeScript, Next.js, Tailwind CSS, Redux, HTML5, CSS3")
    y -= 15
    c.drawString(50, y, "• Backend / Scripting: Node.js, Python (basic), REST API, GraphQL")
    y -= 15
    c.drawString(50, y, "• Tools: Git, Webpack, Figma, Jest, Cypress")
    y -= 15
    c.drawString(50, y, "• Soft Skills: Leadership, Communication, Mentoring, Design Collaboration")
    y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "EXPERIENCE")
    y -= 15
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Lead UI Engineer | Enterprise Cloud (2020 - Present | 4 years)")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "- Spearheaded development of enterprise frontend using Next.js and TypeScript.")
    y -= 20
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Frontend Developer | Creative Studio (2018 - 2020 | 2 years)")
    y -= 15
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "- Built responsive web interfaces with React and modern CSS.")

    c.save()

# 5. Generate michael_devops_cloud.txt
def create_michael_txt(filename: str):
    content = """Michael Vance - Senior DevOps & Cloud Infrastructure Engineer
Email: michael.vance@example.com | Phone: +1-555-0177 | Location: Denver, CO

SUMMARY:
DevOps and Cloud Engineer with 5 years of experience automating infrastructure, cloud operations,
and container platforms across AWS. Proficient in Python scripting, Docker, Kubernetes, and CI/CD.

TECHNICAL SKILLS:
- Cloud Platforms: AWS (EC2, S3, EKS, RDS, CloudWatch), GCP
- Containers & Orchestration: Docker, Kubernetes (K8s), Helm
- Automation & Scripting: Python, Bash, Terraform, Ansible
- CI/CD & Tools: Git, GitHub Actions, Jenkins, Linux
- Soft Skills: Troubleshooting, Problem Solving, Cross-functional Teamwork

EXPERIENCE:
Senior Cloud Engineer | CloudScale Networks (2021 - Present | 3 years)
- Managed Kubernetes clusters on AWS supporting high-traffic microservices.
- Wrote Python automation scripts and Terraform modules for automated cloud provisioning.
- Implemented robust CI/CD deployment pipelines reducing release cycles by 60%.

Infrastructure Engineer | DataPeak Inc. (2019 - 2021 | 2 years)
- Administered Linux server fleet, Docker containers, and automated backup workflows.
- Developed Python monitoring scripts for health checks and alert telemetry.

EDUCATION & CERTIFICATIONS:
- B.S. in Computer Information Systems - Colorado State University
- AWS Certified DevOps Engineer - Professional
"""
    with open(filename, "w", encoding="utf-8") as f:
        f.write(content)

create_alex_pdf("test_data/resumes/alex_senior_python.pdf")
create_sarah_docx("test_data/resumes/sarah_mid_python.docx")
create_jordan_txt("test_data/resumes/jordan_junior_python.txt")
create_emily_pdf("test_data/resumes/emily_frontend_lead.pdf")
create_michael_txt("test_data/resumes/michael_devops_cloud.txt")
print("Successfully generated all 5 test resumes in test_data/resumes/")
